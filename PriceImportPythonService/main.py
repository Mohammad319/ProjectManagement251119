from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

import fitz
import pandas as pd
from docx import Document
from fastapi import FastAPI, File, UploadFile


SERVICE_NAME = "PriceImportPythonService"

app = FastAPI(title=SERVICE_NAME)


@app.get("/health")
async def health() -> dict[str, Any]:
    return {
        "ok": True,
        "service": SERVICE_NAME,
    }


@app.post("/extract-text")
async def extract_text(file: UploadFile = File(...)) -> dict[str, Any]:
    file_name = file.filename or "uploaded-file"
    extension = Path(file_name).suffix.lower().lstrip(".")
    file_type = normalize_file_type(extension)
    content = await file.read()
    errors: list[str] = []
    pages: list[dict[str, Any]] = []
    sheets: list[dict[str, Any]] = []

    try:
        if file_type == "pdf":
            pages = extract_pdf_pages(content)
        elif file_type == "docx":
            pages = extract_docx_pages(content)
        elif file_type == "xlsx":
            sheets = extract_xlsx_sheets(content)
        else:
            errors.append(f"Unsupported file type: {extension or 'unknown'}")
    except Exception as exc:
        errors.append(str(exc))

    return {
        "ok": len(errors) == 0,
        "fileName": file_name,
        "fileType": file_type,
        "pages": pages,
        "sheets": sheets,
        "errors": errors,
    }


def normalize_file_type(extension: str) -> str:
    if extension == "pdf":
        return "pdf"
    if extension == "docx":
        return "docx"
    if extension in {"xlsx", "xls"}:
        return "xlsx"
    return "unknown"


def extract_pdf_pages(content: bytes) -> list[dict[str, Any]]:
    pages: list[dict[str, Any]] = []

    with fitz.open(stream=content, filetype="pdf") as document:
        for index, page in enumerate(document, start=1):
            pages.append(
                {
                    "pageNumber": index,
                    "text": page.get_text("text").strip(),
                }
            )

    return pages


def extract_docx_pages(content: bytes) -> list[dict[str, Any]]:
    document = Document(BytesIO(content))
    blocks: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)

    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                blocks.append("\t".join(values))

    return [
        {
            "pageNumber": 1,
            "text": "\n".join(blocks),
        }
    ]


def extract_xlsx_sheets(content: bytes) -> list[dict[str, Any]]:
    engine = "openpyxl" if file_type_from_content(content) == "xlsx" else None
    workbook = pd.read_excel(
        BytesIO(content),
        sheet_name=None,
        header=None,
        dtype=object,
        engine=engine,
    )
    sheets: list[dict[str, Any]] = []

    for sheet_name, frame in workbook.items():
        rows: list[list[str]] = []
        frame = frame.where(pd.notna(frame), None)

        for row in frame.itertuples(index=False, name=None):
            values = [format_cell_value(value) for value in row]
            if any(value != "" for value in values):
                rows.append(values)

        sheets.append(
            {
                "sheetName": sheet_name,
                "rows": rows,
            }
        )

    return sheets


def format_cell_value(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def file_type_from_content(content: bytes) -> str:
    return "xlsx" if content.startswith(b"PK") else "xls"
